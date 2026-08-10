# Low-level OOXML helpers for building the thesis document with python-docx:
# real Word fields (TOC / PAGEREF) and bookmarks so the Table of Contents and the
# List of Figures / List of Tables compute correctly when the reader updates fields in Word.

from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
import itertools

_bookmark_id = itertools.count(1000)


def add_bookmark(paragraph, name):
    bid = str(next(_bookmark_id))
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _field_char(run, kind, instr=None):
    r = run._r
    if kind == "begin":
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), "begin")
        r.append(fld)
    elif kind == "instr":
        t = OxmlElement("w:instrText")
        t.set(qn("xml:space"), "preserve")
        t.text = instr
        r.append(t)
    elif kind == "separate":
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), "separate")
        r.append(fld)
    elif kind == "end":
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), "end")
        r.append(fld)


def add_field(paragraph, instr_text, result_text="", bold=False, italic=False, size=None):
    r1 = paragraph.add_run()
    _field_char(r1, "begin")
    r2 = paragraph.add_run()
    _field_char(r2, "instr", instr_text)
    r3 = paragraph.add_run()
    _field_char(r3, "separate")
    r4 = paragraph.add_run(result_text)
    r4.bold = bold
    r4.italic = italic
    if size:
        r4.font.size = Pt(size)
    r5 = paragraph.add_run()
    _field_char(r5, "end")
    return r4


def add_pageref(paragraph, bookmark_name, size=None):
    return add_field(paragraph, f'PAGEREF {bookmark_name} \\h', "1", size=size)


def add_toc_field(doc, switches='\\o "1-3" \\h \\z \\u'):
    p = doc.add_paragraph()
    r = p.add_run()
    _field_char(r, "begin")
    r2 = p.add_run()
    _field_char(r2, "instr", f"TOC {switches}")
    r3 = p.add_run()
    _field_char(r3, "separate")
    r4 = p.add_run("Right-click here and choose “Update Field” (or press Ctrl+A then F9) to generate the Table of Contents.")
    r4.italic = True
    r5 = p.add_run()
    _field_char(r5, "end")
    return p


def add_dot_leader_entry(doc, label_text, bookmark_name, style="Normal", indent=None):
    """One List-of-Figures / List-of-Tables row: label ......... <page>."""
    p = doc.add_paragraph(style=style)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
    run = p.add_run(label_text + "\t")
    add_pageref(p, bookmark_name)
    return p


def set_page_numbers(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    add_field(p, "PAGE", "1")


def restart_page_numbers(section, start=1, fmt="decimal"):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn("w:pgNumType"))
    if pgNumType is None:
        pgNumType = OxmlElement("w:pgNumType")
        sectPr.append(pgNumType)
    pgNumType.set(qn("w:start"), str(start))
    pgNumType.set(qn("w:fmt"), fmt)


def shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def set_col_widths(table, widths_inches):
    for row in table.rows:
        for idx, w in enumerate(widths_inches):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(w)
    for idx, w in enumerate(widths_inches):
        if idx < len(table.columns):
            table.columns[idx].width = Inches(w)
